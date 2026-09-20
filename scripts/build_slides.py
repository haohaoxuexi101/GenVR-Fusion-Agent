#!/usr/bin/env python3
"""Build the defense slide deck (PPTX) from official_outputs/SLIDES_CONTENT.md."""

from __future__ import annotations

from pathlib import Path
import re

from pptx import Presentation
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "official_outputs" / "SLIDES_CONTENT.md"
OUT = ROOT / "official_outputs" / "OPENxxx_RayLab-GMC_slides.pptx"

SLIDE_RE = re.compile(r"^\*\*第\s*(\d+)\s*页\s*·\s*(.+?)\*\*\s*$")


def parse() -> list[tuple[str, list[tuple[int, str]]]]:
    slides: list[tuple[str, list[tuple[int, str]]]] = []
    current_title = ""
    current_bullets: list[tuple[int, str]] = []
    for raw in SRC.read_text(encoding="utf-8").splitlines():
        m = SLIDE_RE.match(raw)
        if m:
            if current_title or current_bullets:
                slides.append((current_title, current_bullets))
            current_title = m.group(2).strip()
            current_bullets = []
            continue
        stripped = raw.rstrip()
        if stripped.startswith("- "):
            current_bullets.append((0, stripped[2:].strip()))
        elif stripped.startswith("  - "):
            current_bullets.append((1, stripped[4:].strip()))
    if current_title or current_bullets:
        slides.append((current_title, current_bullets))
    return slides


def main() -> None:
    slides = parse()
    prs = Presentation()

    for index, (title, bullets) in enumerate(slides):
        if index == 0:
            slide = prs.slides.add_slide(prs.slide_layouts[0])
            slide.shapes.title.text = title
            if bullets:
                slide.placeholders[1].text = "\n".join(text for _, text in bullets[:2])
        else:
            slide = prs.slides.add_slide(prs.slide_layouts[1])
            slide.shapes.title.text = title
            body = slide.placeholders[1].text_frame
            body.word_wrap = True
            for level, text in bullets:
                para = body.paragraphs[0] if not body.paragraphs[0].runs and not body.paragraphs[0].text else body.add_paragraph()
                para.text = text
                para.level = level
                for run in para.runs:
                    run.font.size = Pt(16 if level == 0 else 13)
    prs.save(OUT)
    print(f"WROTE {OUT} with {len(slides)} slides")


if __name__ == "__main__":
    main()
